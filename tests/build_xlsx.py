import json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
HOST = os.environ.get("API_HOST", "https://astro.example.com")  # set API_HOST for a private table
out = sys.argv[1] if len(sys.argv) > 1 else "/workspace/天文工具明细.xlsx"
wb = Workbook(); ws = wb.active; ws.title = "工具明细"
heads = ["序号", "库名", "领域", "工具名称", "功能描述", "参数列表\n(名称[类型] 必填/可选: 描述)", "参数数量", "真实调用可用状态（20261012）", "API接口"]
ws.append(["天文工具库 · 工具明细表"] + [None] * 8); ws.merge_cells("A1:I1")
ws["A1"].font = Font(bold=True, size=14); ws["A1"].alignment = Alignment(horizontal="center")
ws.append(heads)
fill = PatternFill("solid", fgColor="DDEBF7"); thin = Side(style="thin", color="BBBBBB")
for c in ws[2]: c.font = Font(bold=True); c.fill = fill; c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
n = 0
for sdk in ["astropy", "astroquery", "photutils", "astroplan", "sunpy"]:
    tools = json.load(open(f"{ROOT}/data/{sdk}_tools.json"))
    res = {r["name"]: r for r in json.load(open(f"{ROOT}/results/{sdk}.json"))}
    for t in tools:
        n += 1
        ws.append([n, sdk, t["domain"], t["name"], t["desc"], t["orig_params"], len(t["params"]), res[t["name"]]["status"],
                   f"POST {HOST}/{sdk}/invoke（function_name={t['name']}）"])
for row in ws.iter_rows(min_row=3):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical="top"); c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
for col, w in zip("ABCDEFGHI", [6, 12, 20, 38, 60, 60, 9, 34, 70]): ws.column_dimensions[col].width = w
ws.freeze_panes = "A3"
wb.save(out); print("saved", out, n)
